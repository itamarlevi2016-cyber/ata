"""ייצוא התמלול לקבצים: TXT, SRT, DOCX."""

import io


def fmt_time(seconds: float, srt: bool = False) -> str:
    ms = int(round(seconds * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    if srt:
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
    return f"{h:02d}:{m:02d}:{s:02d}"


def _label(seg: dict, speakers: dict[str, str]) -> str:
    sp = seg.get("speaker")
    return speakers.get(sp, sp) if sp else ""


def to_txt(result: dict, timestamps: bool = True) -> str:
    speakers = result.get("speakers") or {}
    lines = []
    for seg in result["segments"]:
        prefix = f"[{fmt_time(seg['start'])}] " if timestamps else ""
        label = _label(seg, speakers)
        if label:
            prefix += f"{label}: "
        lines.append(prefix + seg["text"])
    return "\n\n".join(lines) + "\n"


def to_srt(result: dict) -> str:
    speakers = result.get("speakers") or {}
    blocks = []
    for i, seg in enumerate(result["segments"], 1):
        label = _label(seg, speakers)
        text = f"{label}: {seg['text']}" if label else seg["text"]
        blocks.append(f"{i}\n{fmt_time(seg['start'], True)} --> {fmt_time(seg['end'], True)}\n{text}\n")
    return "\n".join(blocks)


def to_docx(result: dict, title: str) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def rtl(paragraph):
        p_pr = paragraph._p.get_or_add_pPr()
        bidi = p_pr.makeelement(qn("w:bidi"), {})
        p_pr.append(bidi)

    def rtl_run(run):
        r_pr = run._r.get_or_add_rPr()
        r_pr.append(r_pr.makeelement(qn("w:rtl"), {}))

    speakers = result.get("speakers") or {}
    doc = Document()
    heading = doc.add_heading(title, level=1)
    rtl(heading)
    for seg in result["segments"]:
        p = doc.add_paragraph()
        rtl(p)
        ts = p.add_run(f"[{fmt_time(seg['start'])}] ")
        ts.font.size = Pt(9)
        ts.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        label = _label(seg, speakers)
        if label:
            r = p.add_run(f"{label}: ")
            r.bold = True
            rtl_run(r)
        r = p.add_run(seg["text"])
        rtl_run(r)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
